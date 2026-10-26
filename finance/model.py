"""Modèle financier mensuel (36 mois) d'Ardane : abonnements et achat à vie vendus par l'App Store.

Usage : python finance/model.py [--structure individuelle|societe] [--params chemin.json]

Principes (aucun double comptage) :
- Prix de liste = prix en USD affiché (6,99 $, 59,99 $, 99,99 $). Le client paie le prix TTC.
- Taxes collectées par Apple : jamais comptées dans le chiffre d'affaires ni dans les proceeds.
- Chiffre d'affaires (HT) = prix de liste x facteur HT. Commission Apple = taux x HT, comptée UNE fois
  (déduite des proceeds, jamais ajoutée aux dépenses).
- Proceeds = HT - commission. Remboursements = ligne séparée. Dépenses = frais propres de l'entreprise.
- Bénéfice comptable avant impôt = proceeds nets (CAD) - dépenses. L'impôt n'y figure pas.
"""
import argparse
import csv
import json
import math
import os
from copy import deepcopy

HERE = os.path.dirname(os.path.abspath(__file__))
MONTHS = 36


# ---------------------------------------------------------------- paramètres
def strip(node):
    """Transforme {value, source, status} en value ; ignore les clés qui commencent par _."""
    if isinstance(node, dict):
        if "value" in node:
            return node["value"]
        return {k: strip(v) for k, v in node.items() if not k.startswith("_")}
    return node


def load_params(path=None):
    with open(path or os.path.join(HERE, "params.json"), encoding="utf-8") as f:
        return strip(json.load(f))


# ---------------------------------------------------------------- briques pures
def price_factors(P):
    """Par dollar de prix de liste : (facteur HT, facteur taxe).
    - boutiques TTC : le client paie P, HT = P/(1+t), taxe = P*t/(1+t)
    - boutiques taxe ajoutée : HT = P, taxe = P*t (le client paie P*(1+t))
    - le reste : sans taxe, HT = P."""
    t = P["taxes_apple"]
    s_in, r_in = t["share_inclusive"], t["rate_inclusive"]
    s_ex, r_ex = t["share_exclusive"], t["rate_exclusive"]
    s_no = 1.0 - s_in - s_ex
    ht = s_in / (1 + r_in) + s_ex + s_no
    tax = s_in * r_in / (1 + r_in) + s_ex * r_ex
    return ht, tax


def commission_rates(P, mode="standard"):
    """(taux 1re année, taux mature). Variante PME : 15 % partout, sans cumul."""
    c = P["commission"]
    if mode == "small_business":
        return c["small_business_rate"], c["small_business_rate"]
    return c["first_rate"], c["mature_rate"]


def downloads(sc, m):
    """Téléchargements du mois m (1 = premier mois). Croissance composée, plafond, mois d'inactivité."""
    start, n = sc["inactive_start_month"], sc["inactive_months"]
    if n > 0 and start <= m < start + n:
        return 0.0
    return min(sc["downloads_m1"] * (1 + sc["growth_monthly"]) ** (m - 1), sc["downloads_cap"])


def marketing(sc, m):
    start, n = sc["inactive_start_month"], sc["inactive_months"]
    if n > 0 and start <= m < start + n:
        return 0.0
    return sc["marketing_monthly_cad"]


def refund_cost(proceeds, commission, rate, commission_kept):
    """Coût des remboursements pour le développeur.
    La part remboursée r des ventes : le développeur rend le prix HT. Il avait reçu HT - commission.
    Si Apple garde sa commission : perte = r*HT = r*(proceeds + commission) ; sinon : r*proceeds."""
    return rate * (proceeds + (commission if commission_kept else 0.0))


def recurring_expenses_annual(P, structure, scenario):
    """Dépenses récurrentes annuelles en CAD (sans création)."""
    d = P["depenses"]
    s = d[structure]
    fx = P["change"]["usd_cad"]
    sc = P["scenarios"][scenario]
    return (s["accountant_annual_cad"] + s["registry_annual_cad"] + d["apple_developer_usd_per_year"] * fx
            + d["insurance_annual_cad"] + d["professional_fees_annual_cad"]
            + 12 * d["hosting_api_ai_monthly_cad"] + 12 * sc["marketing_monthly_cad"])


def expenses_month(P, structure, sc, m):
    """Dépenses du mois m en CAD.
    Création au mois 1 ; registre aux mois 13 et 25 ; comptable lissé par mois ;
    programme développeur Apple au premier renouvellement puis chaque 12 mois."""
    d = P["depenses"]
    s = d[structure]
    fx = P["change"]["usd_cad"]
    out = {"creation": 0.0, "registre": 0.0, "comptable": 0.0, "apple_developer": 0.0,
           "hebergement_api_ia": d["hosting_api_ai_monthly_cad"], "marketing": marketing(sc, m),
           "frais_professionnels": d["professional_fees_annual_cad"] / 12,
           "assurances": d["insurance_annual_cad"] / 12}
    if m == 1:
        out["creation"] = s["creation_cad"]
    if m > 1 and (m - 1) % 12 == 0:
        out["registre"] = s["registry_annual_cad"]
    out["comptable"] = s["accountant_annual_cad"] / 12
    first = d["apple_developer_first_charge_month"]
    if m >= first and (m - first) % 12 == 0:
        out["apple_developer"] = d["apple_developer_usd_per_year"] * fx
    return out


# ---------------------------------------------------------------- simulation
def simulate(P, scenario, structure="individuelle", mode="standard", months=MONTHS):
    """Retourne une liste de dictionnaires (un par mois)."""
    sc = P["scenarios"][scenario]
    pr = P["prix_usd"]
    fx = P["change"]["usd_cad"]
    ht_f, tax_f = price_factors(P)
    rate_first, rate_mature = commission_rates(P, mode)
    mature_after = P["commission"]["mature_after_paid_months"]
    kept = P["remboursements"]["commission_kept_by_apple"]
    ref = sc["refund_rate"]
    churn, renew = sc["churn_monthly"], sc["annual_renewal"]
    w = sc["welcome_share"]

    # Nouveaux payants par mois et par type (cohortes)
    dl = {m: downloads(sc, m) for m in range(1, months + 1)}
    new_m = {m: dl[m] * sc["conversion"] * sc["mix_monthly"] for m in dl}
    new_a = {m: dl[m] * sc["conversion"] * sc["mix_annual"] for m in dl}
    new_l = {m: dl[m] * sc["conversion"] * sc["mix_lifetime"] for m in dl}

    rows = []
    cum_profit = 0.0
    cash = P["paiement"]["opening_cash_cad"]
    net_cad_hist = {}
    for t in range(1, months + 1):
        list_first = 0.0   # ventes au prix de liste soumises au taux 1re année
        list_mature = 0.0  # ventes au prix de liste soumises au taux mature
        pay_m = pay_a = pay_l = 0.0
        resil_m = renouv_a = perdu_a = 0.0
        # Abonnés mensuels : la cohorte c paie son k-ième mois (k = t-c+1)
        for c in range(1, t + 1):
            k = t - c + 1
            alive = new_m[c] * (1 - churn) ** (k - 1)  # payants du mois k
            if k >= 2:
                resil_m += new_m[c] * (1 - churn) ** (k - 2) * churn  # n'ont pas payé le mois k
            price = (w * pr["welcome_monthly"] + (1 - w) * pr["monthly"]) if k <= pr["welcome_months"] else pr["monthly"]
            amount = alive * price
            pay_m += alive
            if k <= mature_after:
                list_first += amount
            else:
                list_mature += amount
        # Abonnés annuels : paiement à l'inscription puis chaque 12 mois
        active_a = 0.0
        for c in range(1, t + 1):
            j = (t - c) // 12  # nombre de renouvellements déjà passés
            active_a += new_a[c] * renew ** j
            if (t - c) % 12 == 0:
                n = new_a[c] * renew ** j
                pay_a += n
                if j == 0:
                    list_first += n * pr["annual"]
                else:
                    list_mature += n * pr["annual"]
                    renouv_a += n
                    perdu_a += new_a[c] * renew ** (j - 1) * (1 - renew)
        # Achats à vie : toujours au taux 1re année (pas d'abonnement)
        pay_l = new_l[t]
        list_first += pay_l * pr["lifetime"]

        list_total = list_first + list_mature
        ht = list_total * ht_f
        taxes = list_total * tax_f
        ttc = ht + taxes
        commission = list_first * ht_f * rate_first + list_mature * ht_f * rate_mature
        proceeds = ht - commission
        refunds = refund_cost(proceeds, commission, ref, kept)
        net_usd = proceeds - refunds
        withholding = net_usd * P["paiement"]["us_withholding_rate"]
        net_usd_after = net_usd - withholding
        net_cad = net_usd_after * fx

        exp = expenses_month(P, structure, sc, t)
        exp_total = sum(exp.values())
        profit = net_cad - exp_total
        cum_profit += profit
        net_cad_hist[t] = net_cad
        cash_in = net_cad_hist.get(t - P["paiement"]["lag_months"], 0.0)
        cash += cash_in - exp_total

        rows.append({
            "mois": t,
            "telechargements": dl[t],
            "nouveaux_mensuels": new_m[t], "nouveaux_annuels": new_a[t], "nouveaux_a_vie": new_l[t],
            "payants_mensuels": pay_m, "resilies_mensuels": resil_m,
            "abonnes_annuels_actifs": active_a, "annuels_renouveles": renouv_a, "annuels_perdus": perdu_a,
            "paiements_annuels": pay_a,
            "ventes_ttc_client_usd": ttc,
            "taxes_collectees_apple_usd": taxes,
            "chiffre_affaires_ht_usd": ht,
            "commission_apple_usd": commission,
            "proceeds_usd": proceeds,
            "remboursements_usd": refunds,
            "retenue_usa_usd": withholding,
            "revenus_nets_usd": net_usd_after,
            "revenus_nets_cad": net_cad,
            "depenses_cad": exp_total,
            "depenses_detail": exp,
            "benefice_avant_impot_cad": profit,
            "benefice_cumule_cad": cum_profit,
            "encaissement_apple_cad": cash_in,
            "tresorerie_cad": cash,
        })
    return rows


# ---------------------------------------------------------------- analyses
def break_even(rows):
    """(mois à partir duquel le bénéfice cumulé reste >= 0 ; premier mois à bénéfice mensuel > 0), None si jamais."""
    neg = [r["mois"] for r in rows if r["benefice_cumule_cad"] < 0]
    if not neg:
        cum_m = 1
    elif neg[-1] == rows[-1]["mois"]:
        cum_m = None  # toujours négatif à la fin de l'horizon
    else:
        cum_m = neg[-1] + 1  # dernier mois négatif + 1 : le cumul reste ensuite >= 0
    op_m = next((r["mois"] for r in rows if r["benefice_avant_impot_cad"] > 0), None)
    return cum_m, op_m


def cash_break_even(rows):
    """Mois à partir duquel la trésorerie reste >= 0 (None si négative à la fin)."""
    neg = [r["mois"] for r in rows if r["tresorerie_cad"] < 0]
    if not neg:
        return 1
    return None if neg[-1] == rows[-1]["mois"] else neg[-1] + 1


def yearly(rows):
    """Agrège par année (1 à 3)."""
    out = []
    for y in range(1, MONTHS // 12 + 1):
        r = rows[(y - 1) * 12: y * 12]
        s = lambda k: sum(x[k] for x in r)
        out.append({
            "annee": y,
            "telechargements": s("telechargements"),
            "ventes_ttc_client_usd": s("ventes_ttc_client_usd"),
            "taxes_collectees_apple_usd": s("taxes_collectees_apple_usd"),
            "chiffre_affaires_ht_usd": s("chiffre_affaires_ht_usd"),
            "commission_apple_usd": s("commission_apple_usd"),
            "remboursements_usd": s("remboursements_usd"),
            "revenus_nets_usd": s("revenus_nets_usd"),
            "revenus_nets_cad": s("revenus_nets_cad"),
            "depenses_cad": s("depenses_cad"),
            "benefice_avant_impot_cad": s("benefice_avant_impot_cad"),
            "tresorerie_fin_cad": r[-1]["tresorerie_cad"],
            "abonnes_actifs_fin": r[-1]["payants_mensuels"] + r[-1]["abonnes_annuels_actifs"],
        })
    return out


def net_per_subscriber_year_usd(P, scenario, rate):
    """Revenus nets USD par abonné actif et par an, régime permanent (hors offre de bienvenue et achats à vie).
    Mix mensuel/annuel du scénario renormalisé. Remboursements inclus comme dans refund_cost :
    net = HT x (1 - c) x (1 - r) - [commission gardée] r x c x HT."""
    sc = P["scenarios"][scenario]
    pr = P["prix_usd"]
    ht_f, _ = price_factors(P)
    share_a = sc["mix_annual"] / (sc["mix_monthly"] + sc["mix_annual"])
    list_year = (1 - share_a) * 12 * pr["monthly"] + share_a * pr["annual"]
    ht = list_year * ht_f
    r = sc["refund_rate"]
    proceeds = ht * (1 - rate)
    commission = ht * rate
    return proceeds - refund_cost(proceeds, commission, r, P["remboursements"]["commission_kept_by_apple"])


def subs_needed(P, scenario, target_cad, structure="individuelle", rate=None):
    """Abonnés actifs (mensuels + annuels) pour atteindre un bénéfice comptable annuel avant impôt.
    nombre = (cible + dépenses récurrentes annuelles) / revenus nets CAD par abonné et par an."""
    if rate is None:
        rate = P["commission"]["mature_rate"]
    fx = P["change"]["usd_cad"]
    per_sub = net_per_subscriber_year_usd(P, scenario, rate) * (1 - P["paiement"]["us_withholding_rate"]) * fx
    need = recurring_expenses_annual(P, structure, scenario) + target_cad
    return math.ceil(need / per_sub)


# ---------------------------------------------------------------- sorties
def _fmt(v):
    return round(v, 2) if isinstance(v, float) else v


def write_monthly(rows, path):
    keys = [k for k in rows[0] if k != "depenses_detail"]
    det = list(rows[0]["depenses_detail"])
    with open(path, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f, delimiter=";")
        w.writerow(keys + ["dep_" + d for d in det])
        for r in rows:
            w.writerow([_fmt(r[k]) for k in keys] + [_fmt(r["depenses_detail"][d]) for d in det])


def run_all(P, structure="individuelle", out_dir=None):
    out_dir = out_dir or os.path.join(HERE, "out")
    os.makedirs(out_dir, exist_ok=True)
    summary, needs, results = [], [], {}
    for name in P["scenarios"]:
        for variant, mode in (("", "standard"), ("_pme", "small_business")):
            rows = simulate(P, name, structure, mode)
            results[name + variant] = rows
            write_monthly(rows, os.path.join(out_dir, name + variant + ".csv"))
            be_cum, be_op = break_even(rows)
            for y in yearly(rows):
                summary.append({"scenario": name + variant, **{k: _fmt(v) for k, v in y.items()},
                                "seuil_rentabilite_cumule_mois": be_cum, "premier_mois_benefice_mensuel": be_op,
                                "tresorerie_positive_des_mois": cash_break_even(rows)})
        for target in P["objectifs_benefice_cad"]:
            needs.append({"scenario": name, "objectif_benefice_avant_impot_cad": target,
                          "abonnes_necessaires_commission_15": subs_needed(P, name, target, structure, P["commission"]["mature_rate"]),
                          "abonnes_necessaires_commission_30": subs_needed(P, name, target, structure, P["commission"]["first_rate"])})
    for fname, data in (("resume.csv", summary), ("abonnes_necessaires.csv", needs)):
        with open(os.path.join(out_dir, fname), "w", newline="", encoding="utf-8-sig") as f:
            w = csv.DictWriter(f, fieldnames=list(data[0]), delimiter=";")
            w.writeheader()
            w.writerows(data)
    # XLSX facultatif
    try:
        import openpyxl  # noqa: F401
        _write_xlsx(P, results, summary, needs, os.path.join(out_dir, "modele.xlsx"))
    except ImportError:
        pass
    return results, summary, needs


def _write_xlsx(P, results, summary, needs, path):
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "hypotheses"
    ws.append(["chemin", "valeur", "source", "statut"])
    def walk(node, p=""):
        for k, v in node.items():
            if k.startswith("_"):
                continue
            if isinstance(v, dict) and "value" in v:
                ws.append([p + k, json.dumps(v["value"], ensure_ascii=False), v.get("source"), v.get("status")])
            elif isinstance(v, dict):
                walk(v, p + k + ".")
    with open(os.path.join(HERE, "params.json"), encoding="utf-8") as f:
        walk(json.load(f))
    for name, rows in results.items():
        s = wb.create_sheet(name)
        keys = [k for k in rows[0] if k != "depenses_detail"]
        s.append(keys)
        for r in rows:
            s.append([_fmt(r[k]) for k in keys])
    for title, data in (("resume", summary), ("abonnes_necessaires", needs)):
        s = wb.create_sheet(title)
        s.append(list(data[0]))
        for d in data:
            s.append(list(d.values()))
    wb.save(path)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--structure", default="individuelle", choices=["individuelle", "societe"])
    ap.add_argument("--params")
    a = ap.parse_args()
    P = load_params(a.params)
    results, summary, needs = run_all(P, a.structure)
    for s in summary:
        print(s)
    for n in needs:
        print(n)


if __name__ == "__main__":
    main()
